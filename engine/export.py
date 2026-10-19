"""Export in Excel (xlsx) e CSV."""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .models import Risultato

INTESTAZIONI = ["Data", "Data valuta", "Descrizione", "Categoria", "Entrate", "Uscite", "Importo"]


def to_csv(ris: Risultato, separatore: str = ";") -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=separatore, lineterminator="\n")
    w.writerow(INTESTAZIONI)
    for m in ris.movimenti:
        w.writerow([
            m.data.strftime("%d/%m/%Y"),
            m.data_valuta.strftime("%d/%m/%Y") if m.data_valuta else "",
            m.descrizione,
            m.categoria,
            _it(m.importo) if m.importo > 0 else "",
            _it(-m.importo) if m.importo < 0 else "",
            _it(m.importo),
        ])
    return buf.getvalue()


def _it(v: Decimal) -> str:
    return f"{v:.2f}".replace(".", ",")


def to_xlsx(ris: Risultato) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimenti"

    blu = PatternFill("solid", fgColor="1F3A5F")
    bianco = Font(color="FFFFFF", bold=True)
    bordo = Border(bottom=Side(style="thin", color="D0D7E2"))

    # intestazione
    ws.append(INTESTAZIONI)
    for c in ws[1]:
        c.fill = blu
        c.font = bianco
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for m in ris.movimenti:
        ws.append([
            m.data,
            m.data_valuta,
            m.descrizione,
            m.categoria,
            float(m.importo) if m.importo > 0 else None,
            float(-m.importo) if m.importo < 0 else None,
            float(m.importo),
        ])

    n = len(ris.movimenti)
    for r in range(2, n + 2):
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
        ws.cell(r, 2).number_format = "DD/MM/YYYY"
        for col in (5, 6, 7):
            ws.cell(r, col).number_format = '#,##0.00 €;[Red]-#,##0.00 €'
        for col in range(1, 8):
            ws.cell(r, col).border = bordo

    # riga totali con formule vere
    if n:
        tr = n + 2
        ws.cell(tr, 3, "TOTALI").font = Font(bold=True)
        for col, lettera in ((5, "E"), (6, "F"), (7, "G")):
            c = ws.cell(tr, col, f"=SUM({lettera}2:{lettera}{n + 1})")
            c.font = Font(bold=True)
            c.number_format = '#,##0.00 €;[Red]-#,##0.00 €'

    for i, larg in enumerate([12, 12, 60, 28, 14, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.auto_filter.ref = f"A1:G{max(n + 1, 1)}"

    # foglio riepilogo
    rs = wb.create_sheet("Riepilogo")
    rs.append(["Banca", ris.banca])
    rs.append(["Pagine", ris.pagine])
    rs.append(["Movimenti", n])
    q = ris.quadratura
    if q:
        rs.append(["Saldo iniziale", float(q.saldo_iniziale) if q.saldo_iniziale is not None else "n/d"])
        rs.append(["Totale entrate", float(ris.entrate)])
        rs.append(["Totale uscite", float(ris.uscite)])
        rs.append(["Saldo finale dichiarato", float(q.saldo_finale) if q.saldo_finale is not None else "n/d"])
        rs.append(["Saldo ricalcolato", float(q.saldo_calcolato) if q.saldo_calcolato is not None else "n/d"])
        rs.append(["Quadratura", "OK" if q.quadra else ("NON VERIFICABILE" if q.quadra is None else "DIFFERENZA")])
        rs.append(["Nota", q.messaggio])
    rs.append([])
    rs.append(["Categoria", "Entrate", "Uscite", "Netto"])
    for c in rs[rs.max_row]:
        c.font = Font(bold=True)
    per_cat: dict[str, list[Decimal]] = defaultdict(lambda: [Decimal(0), Decimal(0)])
    for m in ris.movimenti:
        if m.importo > 0:
            per_cat[m.categoria][0] += m.importo
        else:
            per_cat[m.categoria][1] += m.importo
    for cat, (e, u) in sorted(per_cat.items(), key=lambda kv: kv[1][1]):
        rs.append([cat, float(e), float(u), float(e + u)])
    for row in rs.iter_rows(min_row=1, max_row=rs.max_row):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = '#,##0.00 €;[Red]-#,##0.00 €'
    rs.column_dimensions["A"].width = 30
    for L in "BCD":
        rs.column_dimensions[L].width = 16

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
